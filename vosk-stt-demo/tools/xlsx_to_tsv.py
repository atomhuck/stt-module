#!/usr/bin/env python3
"""Конвертирует банк вопросов (xlsx) в src/main/resources/interview/questions.tsv.

Использование:
    pip install openpyxl
    python tools/xlsx_to_tsv.py java_backend_ai_interview_question_bank.xlsx

Колонки xlsx: Уровень | Тема | Вопрос | Эталонный ответ | Критерии/заметки | Уточняющий вопрос |
              Уточняющие вопросы (нумерованный список) | Кол-во.
Формат TSV: level \t topic \t question \t reference \t notes \t followup1 || followup2 || ...
"""
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

src = Path(sys.argv[1])
dst = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).resolve().parent.parent / "src/main/resources/interview/questions.tsv"

def clean(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()

ws = load_workbook(src, read_only=True).active
rows = list(ws.iter_rows(values_only=True))[1:]
out = []
for r in rows:
    if not r or not r[2]:
        continue
    level, topic, question, ref, notes, fu_first, fu_all = (r + (None,) * 8)[:7]
    followups = [clean(re.sub(r"^\s*\d+[.)]\s*", "", line)) for line in str(fu_all or "").splitlines() if line.strip()]
    first = clean(fu_first)
    if first and first not in followups:
        followups.insert(0, first)
    out.append("\t".join([clean(level), clean(topic), clean(question), clean(ref), clean(notes), " || ".join(followups)]))

dst.parent.mkdir(parents=True, exist_ok=True)
dst.write_text("\n".join(out) + "\n", encoding="utf-8")
print(f"{len(out)} вопросов -> {dst}")
